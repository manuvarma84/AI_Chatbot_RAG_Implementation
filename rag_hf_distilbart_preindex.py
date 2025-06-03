import os
import json
from datetime import datetime
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import FAISS
from langchain_community.docstore.in_memory import InMemoryDocstore
from langchain_huggingface import HuggingFaceEmbeddings
from langchain.schema import Document
from docx import Document as DocxDocument
import faiss

# Path to metadata storage
METADATA_FILE = <<provide the path of the index_metadata.json here>>
VECTORSTORE_PATH = <<provide the path of vectorstore here>>

# Initialize embeddings
embedding_model_name = "sentence-transformers/all-MiniLM-L6-v2"
embeddings = HuggingFaceEmbeddings(model_name=embedding_model_name)

# Load existing metadata
def load_metadata():
    if os.path.exists(METADATA_FILE):
        with open(METADATA_FILE, "r") as f:
            return json.load(f)
    return {}

# Save metadata
def save_metadata(metadata):
    with open(METADATA_FILE, "w") as f:
        json.dump(metadata, f, indent=4)

# Load and preprocess documents
def load_documents(file_directory, metadata):
    documents = []
    updated_metadata = metadata.copy()
    for filename in os.listdir(file_directory):
        if filename.endswith(".docx"):
            file_path = os.path.join(file_directory, filename)
            file_mtime = os.path.getmtime(file_path)  # Get modification time

            # Check if file is new or updated
            if file_path not in metadata or metadata[file_path]["mtime"] < file_mtime:
                print(f"Indexing new/updated file: {file_path}")
                doc = DocxDocument(file_path)
                text = "\n".join([para.text for para in doc.paragraphs])
                langchain_doc = Document(page_content=text, metadata={"source": file_path})
                documents.append(langchain_doc)

                # Update metadata
                updated_metadata[file_path] = {"mtime": file_mtime}

    return documents, updated_metadata

# Update vectorstore with user confirmation for overwriting
def update_vectorstore(file_directory, vectorstore, metadata):
    # Load new or updated documents
    documents, updated_metadata = load_documents(file_directory, metadata)

    if documents:
        # Split documents into smaller chunks
        text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=100)
        texts = text_splitter.split_documents(documents)

        # Add new embeddings to the vectorstore
        vectorstore.add_documents(texts)
        print(f"Indexed {len(texts)} document chunks.")

        # Return the updated metadata and indicate that the vectorstore needs saving
        return updated_metadata, True
    else:
        print("No new or updated documents to index.")
        
        # Check if the user wants to overwrite even if there are no changes
        user_choice = input("No new documents detected. Do you still want to overwrite the vectorstore? (yes/no): ").strip().lower()
        if user_choice == "yes":
            return updated_metadata, True
        else:
            return updated_metadata, False

# Main Function
if __name__ == "__main__":
    file_directory = <<provide the path of documents folder here>>

    # Load metadata
    metadata = load_metadata()

    # Load or create vectorstore
    if os.path.exists(os.path.join(VECTORSTORE_PATH, "index.faiss")):
        # Load the existing vectorstore
        vectorstore = FAISS.load_local(VECTORSTORE_PATH, embeddings, allow_dangerous_deserialization=True)
        print("Loaded existing vectorstore.")
    else:
        # Initialize a new FAISS index
        dim = len(embeddings.embed_query("test"))  # Get embedding dimension
        index = faiss.IndexFlatL2(dim)  # Create an empty FAISS index
        vectorstore = FAISS(
            index=index,
            embedding_function=embeddings.embed_query,
            docstore=InMemoryDocstore({}),
            index_to_docstore_id={}
        )
        print("Created a new vectorstore.")

    # Update vectorstore
    updated_metadata, needs_saving = update_vectorstore(file_directory, vectorstore, metadata)

    if needs_saving:
        # Save updated vectorstore and metadata
        vectorstore.save_local(VECTORSTORE_PATH)
        save_metadata(updated_metadata)
        print("Updated vectorstore saved.")
